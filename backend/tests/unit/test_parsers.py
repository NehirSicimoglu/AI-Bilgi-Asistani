"""Parser testleri (PDF/DOCX/TXT/MD in-memory üretilir)."""

from __future__ import annotations

import io

import fitz
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.exceptions import UnsupportedFileTypeError
from app.ingestion.parsers import get_parser, parse_document


def _make_pdf(pages_text: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages_text:
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _make_docx(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # varsayılan boş sayfayı kaldır
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_pdf_parser_keeps_pages() -> None:
    data = _make_pdf(["Birinci sayfa metni", "Ikinci sayfa metni"])
    parsed = parse_document(data, "x.pdf", ".pdf")
    assert len(parsed.pages) == 2
    assert parsed.pages[0].page == 1
    assert "Birinci" in parsed.pages[0].text


def test_docx_parser() -> None:
    data = _make_docx(["Merhaba dunya", "Ikinci paragraf"])
    parsed = parse_document(data, "x.docx", ".docx")
    assert parsed.pages[0].page is None
    assert "Merhaba" in parsed.full_text()


def test_txt_parser_utf8() -> None:
    parsed = parse_document("çğıöşü içerik".encode(), "x.txt", ".txt")
    assert "içerik" in parsed.full_text()


def test_txt_parser_non_utf8() -> None:
    data = "kağıt".encode("iso-8859-9")  # Türkçe latin-5
    parsed = parse_document(data, "x.txt", ".txt")
    assert parsed.full_text()  # replace ile de olsa çözülür


def test_markdown_parser() -> None:
    parsed = parse_document(b"# Baslik\n\nParagraf", "x.md", ".md")
    assert "Baslik" in parsed.full_text()


def test_excel_parser_labeled_rows() -> None:
    data = _make_xlsx(
        {
            "Satışlar": [
                ["Ürün", "Adet", "Fiyat"],
                ["Laptop", 12, 25000],
                ["Telefon", 40, 15000],
            ]
        }
    )
    parsed = parse_document(data, "x.xlsx", ".xlsx")
    text = parsed.full_text()
    # Sayfa adı korunur
    assert "Satışlar" in text
    # Her satır başlık-değer çiftleriyle yazılır (chunk'a dayanıklı)
    assert "Ürün: Laptop | Adet: 12 | Fiyat: 25000" in text
    assert "Ürün: Telefon | Adet: 40 | Fiyat: 15000" in text


def test_excel_parser_multiple_sheets() -> None:
    data = _make_xlsx(
        {
            "Q1": [["Ay", "Ciro"], ["Ocak", 100]],
            "Q2": [["Ay", "Ciro"], ["Nisan", 200]],
        }
    )
    parsed = parse_document(data, "x.xlsx", ".xlsx")
    assert len(parsed.pages) == 2
    full = parsed.full_text()
    assert "Ay: Ocak | Ciro: 100" in full
    assert "Ay: Nisan | Ciro: 200" in full


def test_unsupported_extension() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        get_parser(".csv")
