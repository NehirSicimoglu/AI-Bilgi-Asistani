"""Doküman parser'ları — PDF, DOCX, XLSX, TXT, Markdown.

Her parser ham baytları metne çevirir. PDF için sayfa bilgisi korunur (citation'da
sayfa numarası göstermek için); diğer formatlarda sayfa kavramı yoktur (page=None).
Parser'lar swap edilebilir sağlayıcı değil, düz bileşenlerdir; bu yüzden interface
yükü minimaldir (basit bir taban sınıf + factory).
"""

from __future__ import annotations

from abc import ABC, abstractmethod

import chardet
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pydantic import BaseModel

from app.exceptions import AppError, CorruptFileError, UnsupportedFileTypeError


class ParsedPage(BaseModel):
    page: int | None
    text: str


class ParsedDocument(BaseModel):
    """Bir dokümandan çıkarılan metin — sayfa segmentleri halinde."""

    pages: list[ParsedPage]

    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())


class DocumentParser(ABC):
    """Tüm parser'ların ortak arayüzü."""

    @abstractmethod
    def parse(self, data: bytes) -> ParsedDocument: ...


class PdfParser(DocumentParser):
    def parse(self, data: bytes) -> ParsedDocument:
        pages: list[ParsedPage] = []
        with fitz.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc, start=1):
                text = page.get_text("text").strip()
                if text:
                    pages.append(ParsedPage(page=i, text=text))
        if not pages:
            pages.append(ParsedPage(page=None, text=""))
        return ParsedDocument(pages=pages)


class DocxParser(DocumentParser):
    def parse(self, data: bytes) -> ParsedDocument:
        import io

        doc = DocxDocument(io.BytesIO(data))
        parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        # Tablolardaki metni de dahil et
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return ParsedDocument(pages=[ParsedPage(page=None, text="\n".join(parts))])


class ExcelParser(DocumentParser):
    """XLSX için — her satır kendi başlık-değer çiftleriyle metne çevrilir.

    Tablodaki ilk satır başlık kabul edilir; sonraki her satır
    ``Başlık: değer | Başlık: değer`` biçiminde tek satır olarak yazılır. Böylece
    chunking satırları böldüğünde bile her kayıt kendi kendine anlamlı kalır ve
    retrieval sağlamlaşır. Her sayfa (worksheet) ayrı bir ParsedPage olur; sayfa
    kavramı olmadığından ``page=None`` verilir, sayfa adı metnin başına yazılır.
    """

    def parse(self, data: bytes) -> ParsedDocument:
        import io

        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            pages: list[ParsedPage] = []
            for ws in wb.worksheets:
                lines = self._sheet_to_lines(ws)
                if not lines:
                    continue
                text = f"Sayfa: {ws.title}\n\n" + "\n".join(lines)
                pages.append(ParsedPage(page=None, text=text))
        finally:
            wb.close()

        if not pages:
            pages.append(ParsedPage(page=None, text=""))
        return ParsedDocument(pages=pages)

    @staticmethod
    def _cell_to_str(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "Evet" if value else "Hayır"
        return str(value).strip()

    def _sheet_to_lines(self, ws) -> list[str]:
        """İlk dolu satırı başlık kabul edip her veri satırını tek satıra çevirir."""
        header: list[str] | None = None
        out: list[str] = []
        for raw in ws.iter_rows(values_only=True):
            values = [self._cell_to_str(v) for v in raw]
            if not any(values):
                continue  # tamamen boş satırı atla
            if header is None:
                header = values
                continue
            pairs = [
                f"{col}: {val}"
                for col, val in zip(header, values, strict=False)
                if val and col
            ]
            # Başlıkla eşleşmeyen fazladan sütunları da kaybetme
            pairs.extend(v for v in values[len(header) :] if v)
            if pairs:
                out.append(" | ".join(pairs))
        return out


class TextParser(DocumentParser):
    """TXT ve Markdown için — encoding otomatik tespit edilir."""

    def parse(self, data: bytes) -> ParsedDocument:
        text = _decode_bytes(data)
        return ParsedDocument(pages=[ParsedPage(page=None, text=text)])


def _decode_bytes(data: bytes) -> str:
    if not data:
        return ""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        detected = chardet.detect(data)
        encoding = detected.get("encoding") or "latin-1"
        return data.decode(encoding, errors="replace")


_PARSERS: dict[str, DocumentParser] = {
    ".pdf": PdfParser(),
    ".docx": DocxParser(),
    ".xlsx": ExcelParser(),
    ".txt": TextParser(),
    ".md": TextParser(),
}


def get_parser(extension: str) -> DocumentParser:
    """Uzantıya göre uygun parser'ı döndürür."""
    parser = _PARSERS.get(extension.lower())
    if parser is None:
        raise UnsupportedFileTypeError(
            f"Desteklenmeyen dosya türü: {extension}. "
            f"Desteklenen: {', '.join(sorted(_PARSERS))}"
        )
    return parser


def parse_document(data: bytes, filename: str, extension: str) -> ParsedDocument:
    """Ham baytları metne çevirir; bozuk dosyayı 422'ye çevirir.

    Parser kütüphaneleri (PyMuPDF/python-docx/openpyxl) bozuk girdide kendi
    exception'larını fırlatır. Bunlar sarmalanmazsa handler'a "beklenmeyen hata"
    olarak düşer ve kullanıcı 500 görür; oysa yarım inmiş ya da uzantısı
    değiştirilmiş dosya sıradan bir kullanıcı hatasıdır.
    """
    parser = get_parser(extension)
    try:
        return parser.parse(data)
    except AppError:
        raise
    except Exception as exc:
        raise CorruptFileError(
            f"Dosya okunamadı, içerik bozuk veya biçim uyumsuz: {filename}"
        ) from exc
